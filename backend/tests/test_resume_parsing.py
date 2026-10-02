"""Tests for resume parsing and vector storage pipeline."""
from __future__ import annotations

import io
from pathlib import Path
import shutil
import tempfile
import uuid

import pytest
from httpx import ASGITransport, AsyncClient

from app.core.config import settings
from app.core.security import create_access_token
from app.main import app
from app.models.resume import Resume
from app.models.user import User
from app.services.resume_parser import extract_raw_text, resume_parser
from app.services.vector_store import vector_store
from tests.conftest import TestingSessionLocal


@pytest.fixture(autouse=True)
def setup_test_directories():
    """Isolate uploads and ChromaDB persistent directory for testing."""
    temp_upload = tempfile.mkdtemp()
    temp_chroma = tempfile.mkdtemp()

    orig_upload_dir = settings.RESUME_UPLOAD_DIR
    orig_chroma_dir = settings.CHROMA_PERSIST_DIR

    settings.RESUME_UPLOAD_DIR = Path(temp_upload)
    settings.CHROMA_PERSIST_DIR = Path(temp_chroma)
    vector_store.persist_dir = Path(temp_chroma)

    yield

    settings.RESUME_UPLOAD_DIR = orig_upload_dir
    settings.CHROMA_PERSIST_DIR = orig_chroma_dir
    vector_store.persist_dir = orig_chroma_dir
    shutil.rmtree(temp_upload, ignore_errors=True)
    shutil.rmtree(temp_chroma, ignore_errors=True)


async def create_user(email: str) -> tuple[User, str]:
    user_id = uuid.uuid4()
    async with TestingSessionLocal() as session:
        user = User(
            id=user_id,
            google_id=f"google_{user_id}",
            email=email,
            name="Alice Applicant",
            onboarding_completed=False,
            onboarding_step=2,
        )
        session.add(user)
        await session.commit()

    token = create_access_token(data={"sub": str(user_id)})
    return user, token


def create_sample_pdf(file_path: Path) -> None:
    """Generate a real PDF resume using PyMuPDF."""
    import fitz

    doc = fitz.open()
    page = doc.new_page()

    text = """Alex Morgan
alex.morgan@example.com | (555) 123-4567 | San Francisco, CA

SUMMARY
Experienced Software Engineer with a passion for building AI-driven web applications and scalable APIs.

SKILLS
Python, FastAPI, TypeScript, React, PostgreSQL, Docker, PyTorch

EXPERIENCE
Senior Backend Engineer at CloudTech Inc
- Architected and built high-performance microservices using FastAPI and PostgreSQL.
- Reduced API latency by 45% using Redis caching and query optimizations.

Software Developer at StartupSphere
- Developed REST APIs and full-stack features using Python and React.
- Integrated sentence-transformers for search autocomplete.

EDUCATION
B.S. in Computer Science, University of California, Berkeley (2018 - 2022)

PROJECTS
Aptly Job Tracker
- Full-stack AI job matching platform built with FastAPI and React.

CERTIFICATIONS
AWS Certified Solutions Architect - Associate
"""
    p = fitz.Point(50, 50)
    page.insert_text(p, text, fontsize=10)
    doc.save(str(file_path))
    doc.close()


def create_sample_docx(file_path: Path) -> None:
    """Generate a real DOCX resume using python-docx."""
    import docx

    doc = docx.Document()
    doc.add_paragraph("Jordan Taylor")
    doc.add_paragraph("jordan.taylor@example.com | 555-987-6543")
    doc.add_paragraph("SUMMARY")
    doc.add_paragraph("Results-driven Product Engineer specialized in modern distributed cloud systems.")
    doc.add_paragraph("SKILLS")
    doc.add_paragraph("Go, Python, Kubernetes, GraphQL, ChromaDB, Next.js")
    doc.add_paragraph("EXPERIENCE")
    doc.add_paragraph("Lead Engineer at FinTech Labs")
    doc.add_paragraph("- Scaled transaction processing engine to 50k requests per second.")
    doc.add_paragraph("EDUCATION")
    doc.add_paragraph("M.S. in Software Engineering, Stanford University")
    doc.add_paragraph("PROJECTS")
    doc.add_paragraph("Semantic Vector Search Engine")
    doc.add_paragraph("CERTIFICATIONS")
    doc.add_paragraph("Certified Kubernetes Administrator (CKA)")
    doc.save(str(file_path))


@pytest.mark.anyio
async def test_pdf_extraction_and_section_parsing():
    """Verify PyMuPDF extracts text and deterministic parser structures sections."""
    temp_dir = Path(tempfile.mkdtemp())
    pdf_path = temp_dir / "sample_resume.pdf"
    try:
        create_sample_pdf(pdf_path)
        raw_text = extract_raw_text(pdf_path)

        assert "Alex Morgan" in raw_text
        assert "alex.morgan@example.com" in raw_text
        assert "CloudTech Inc" in raw_text

        # Structured deterministic parsing
        parsed = resume_parser.parse(raw_text)
        assert parsed["name"] == "Alex Morgan"
        assert parsed["email"] == "alex.morgan@example.com"
        assert parsed["phone"] is not None
        assert "Python" in parsed["skills"] or any("Python" in s for s in parsed["skills"])
        assert len(parsed["experience"]) >= 1
        assert len(parsed["education"]) >= 1
        assert len(parsed["projects"]) >= 1
        assert len(parsed["certifications"]) >= 1
    finally:
        shutil.rmtree(temp_dir, ignore_errors=True)


@pytest.mark.anyio
async def test_docx_extraction_and_section_parsing():
    """Verify python-docx extracts text and deterministic parser structures sections."""
    temp_dir = Path(tempfile.mkdtemp())
    docx_path = temp_dir / "sample_resume.docx"
    try:
        create_sample_docx(docx_path)
        raw_text = extract_raw_text(docx_path)

        assert "Jordan Taylor" in raw_text
        assert "jordan.taylor@example.com" in raw_text

        parsed = resume_parser.parse(raw_text)
        assert parsed["name"] == "Jordan Taylor"
        assert parsed["email"] == "jordan.taylor@example.com"
        assert len(parsed["skills"]) >= 1
        assert len(parsed["experience"]) >= 1
        assert len(parsed["education"]) >= 1
    finally:
        shutil.rmtree(temp_dir, ignore_errors=True)


@pytest.mark.anyio
async def test_chromadb_vector_creation_and_metadata():
    """Verify ChromaDB stores section chunks with user_id, resume_id, section, chunk_index."""
    user_id = uuid.uuid4()
    resume_id = uuid.uuid4()

    sample_parsed = {
        "summary": "Full stack engineer with 5 years experience.",
        "skills": ["Python", "FastAPI", "React", "PostgreSQL"],
        "experience": ["Senior Backend Engineer at TechCorp (2020 - Present)"],
        "education": ["BS in Computer Science from MIT"],
        "projects": ["Aptly Job Application Tracker"],
        "certifications": ["AWS Certified Developer"],
    }
    raw_text = "Full stack engineer. Skills: Python, React. TechCorp."

    chunks_stored = vector_store.store_resume_chunks(
        resume_id=resume_id,
        user_id=user_id,
        parsed_data=sample_parsed,
        raw_text=raw_text,
    )

    assert chunks_stored >= 5

    # Query ChromaDB collection directly
    results = vector_store.collection.get(
        where={
            "$and": [
                {"user_id": {"$eq": str(user_id)}},
                {"resume_id": {"$eq": str(resume_id)}},
            ]
        }
    )
    assert len(results["ids"]) == chunks_stored
    for meta in results["metadatas"]:
        assert meta["user_id"] == str(user_id)
        assert meta["resume_id"] == str(resume_id)
        assert "section" in meta
        assert "chunk_index" in meta


@pytest.mark.anyio
async def test_parse_resume_api_flow():
    """End-to-end API test: upload -> parse -> verify DB & ChromaDB -> verify GET /current/parsed."""
    user, token = await create_user("resume_pipeline_test@example.com")
    headers = {"Authorization": f"Bearer {token}"}

    # 1. Create PDF resume in bytes
    temp_dir = Path(tempfile.mkdtemp())
    pdf_path = temp_dir / "alex_resume.pdf"
    create_sample_pdf(pdf_path)
    pdf_bytes = pdf_path.read_bytes()
    shutil.rmtree(temp_dir, ignore_errors=True)

    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as ac:
        # Upload
        files = {"file": ("Alex_Morgan_Resume.pdf", io.BytesIO(pdf_bytes), "application/pdf")}
        upload_resp = await ac.post("/api/v1/resumes/upload", headers=headers, files=files)
        assert upload_resp.status_code == 200
        resume_id = upload_resp.json()["id"]

        # Parse
        parse_resp = await ac.post(f"/api/v1/resumes/{resume_id}/parse", headers=headers)
        assert parse_resp.status_code == 200
        parse_data = parse_resp.json()

        assert parse_data["parsed_status"] == "parsed"
        assert parse_data["parsed_at"] is not None
        assert parse_data["parse_error"] is None
        assert parse_data["chunks_stored"] > 0
        assert "Alex Morgan" in parse_data["extracted_text_preview"]
        assert parse_data["parsed_data"]["email"] == "alex.morgan@example.com"

        # Verify DB persisted fields
        async with TestingSessionLocal() as session:
            resume = await session.get(Resume, uuid.UUID(resume_id))
            assert resume is not None
            assert resume.parsed_status == "parsed"
            assert "Alex Morgan" in resume.extracted_text
            assert resume.parsed_data is not None

        # Verify GET /api/v1/resumes/current/parsed
        current_parsed_resp = await ac.get("/api/v1/resumes/current/parsed", headers=headers)
        assert current_parsed_resp.status_code == 200
        current_data = current_parsed_resp.json()
        assert current_data["id"] == resume_id
        assert current_data["parsed_status"] == "parsed"


@pytest.mark.anyio
async def test_user_cannot_access_or_parse_another_users_resume():
    """Security verification: User B cannot parse or access User A's resume."""
    user_a, token_a = await create_user("usera@example.com")
    user_b, token_b = await create_user("userb@example.com")

    headers_a = {"Authorization": f"Bearer {token_a}"}
    headers_b = {"Authorization": f"Bearer {token_b}"}

    temp_dir = Path(tempfile.mkdtemp())
    pdf_path = temp_dir / "usera_resume.pdf"
    create_sample_pdf(pdf_path)
    pdf_bytes = pdf_path.read_bytes()
    shutil.rmtree(temp_dir, ignore_errors=True)

    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as ac:
        # User A uploads resume
        files = {"file": ("UserA_Resume.pdf", io.BytesIO(pdf_bytes), "application/pdf")}
        upload_resp = await ac.post("/api/v1/resumes/upload", headers=headers_a, files=files)
        resume_id_a = upload_resp.json()["id"]

        # User B attempts to parse User A's resume -> should be 403 Forbidden
        hack_resp = await ac.post(f"/api/v1/resumes/{resume_id_a}/parse", headers=headers_b)
        assert hack_resp.status_code == 403
        assert "permission" in hack_resp.json()["detail"].lower() or "access" in hack_resp.json()["detail"].lower()


@pytest.mark.anyio
async def test_parsing_failure_and_retry():
    """Verify error handling on corrupt file, setting failed status and subsequent retry."""
    user, token = await create_user("corrupt_test@example.com")
    headers = {"Authorization": f"Bearer {token}"}

    # Upload empty-like or corrupt file
    corrupt_pdf = b"%PDF-1.4\nCorrupted content with no text stream\n%%EOF"
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as ac:
        files = {"file": ("corrupt.pdf", io.BytesIO(corrupt_pdf), "application/pdf")}
        upload_resp = await ac.post("/api/v1/resumes/upload", headers=headers, files=files)
        assert upload_resp.status_code == 200
        resume_id = upload_resp.json()["id"]

        # Call parse -> should fail safely without crashing
        parse_resp = await ac.post(f"/api/v1/resumes/{resume_id}/parse", headers=headers)
        assert parse_resp.status_code == 200
        data = parse_resp.json()
        assert data["parsed_status"] == "failed"
        assert data["parse_error"] is not None

        # Re-upload valid file (retry flow)
        temp_dir = Path(tempfile.mkdtemp())
        pdf_path = temp_dir / "valid_resume.pdf"
        create_sample_pdf(pdf_path)
        valid_pdf_bytes = pdf_path.read_bytes()
        shutil.rmtree(temp_dir, ignore_errors=True)

        files_valid = {"file": ("valid.pdf", io.BytesIO(valid_pdf_bytes), "application/pdf")}
        reupload_resp = await ac.post("/api/v1/resumes/upload", headers=headers, files=files_valid)
        assert reupload_resp.status_code == 200
        valid_resume_id = reupload_resp.json()["id"]

        # Parse retry
        retry_parse_resp = await ac.post(f"/api/v1/resumes/{valid_resume_id}/parse", headers=headers)
        assert retry_parse_resp.status_code == 200
        retry_data = retry_parse_resp.json()
        assert retry_data["parsed_status"] == "parsed"
        assert retry_data["parse_error"] is None
